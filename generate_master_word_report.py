import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

workspace = r"c:\Users\Muawiyah\Desktop\Jenkins-Based Research Paper Repository"
evidence_dir = os.path.join(workspace, "docs", "evidence")
out_docx = os.path.join(workspace, "docs", "PROJECT_REPORT.docx")

# Styling Colors
COLOR_NAVY = RGBColor(30, 58, 138)       # #1e3a8a
COLOR_SLATE = RGBColor(37, 99, 235)      # #2563eb
COLOR_DARK = RGBColor(30, 41, 59)        # #1e293b
COLOR_GRAY = RGBColor(100, 116, 139)     # #64748b
COLOR_GREEN = RGBColor(22, 101, 52)      # #166534

HEX_NAVY = "1E3A8A"
HEX_LIGHT_BG = "F8FAFC"
HEX_BORDER = "CBD5E1"
HEX_SUCCESS_BG = "DCFCE7"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.name = "Arial"
    if level == 1:
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = COLOR_NAVY
    elif level == 2:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_SLATE
    else:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(2)
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_DARK
    return h

def add_code_snippet(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, HEX_LIGHT_BG)
    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_NAVY}"/>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    run.font.color.rgb = COLOR_DARK
    
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_figure(doc, img_name, fig_num, caption):
    img_path = os.path.join(evidence_dir, img_name)
    if os.path.exists(img_path):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        doc.add_picture(img_path, width=Inches(6.0))
        
        cap_p = doc.add_paragraph()
        cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_p.paragraph_format.space_before = Pt(2)
        cap_p.paragraph_format.space_after = Pt(12)
        r = cap_p.add_run(f"Figure {fig_num}: {caption}")
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.italic = True
        r.font.color.rgb = COLOR_GRAY

def main():
    doc = Document()
    
    # Configure 1-inch margins
    for sec in doc.sections:
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)
        
        # Header / Footer
        header = sec.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Jenkins-Based Research Paper Repository | Master Academic Project Report")
        hrun.font.name = "Arial"
        hrun.font.size = Pt(8)
        hrun.font.color.rgb = COLOR_GRAY
        
        footer = sec.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Candidate: Muawiyah Sahil Narvel (Roll No: 23108A0020)")
        frun.font.name = "Arial"
        frun.font.size = Pt(8.5)
        frun.font.color.rgb = COLOR_GRAY

    # Title Banner
    tp = doc.add_paragraph()
    tp.paragraph_format.space_before = Pt(0)
    tp.paragraph_format.space_after = Pt(2)
    trun = tp.add_run("ACADEMIC PROJECT REPORT:\nJENKINS-BASED RESEARCH PAPER REPOSITORY")
    trun.font.name = "Arial"
    trun.font.size = Pt(18)
    trun.font.bold = True
    trun.font.color.rgb = COLOR_NAVY

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(12)
    srun = sub_p.add_run("Enterprise Software Engineering & DevOps CI/CD Verification Dossier")
    srun.font.name = "Arial"
    srun.font.size = Pt(11)
    srun.font.color.rgb = COLOR_SLATE

    # Candidate Profile Table
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    meta_rows = [
        ("Student Name", "Muawiyah Sahil Narvel"),
        ("Roll Number", "23108A0020"),
        ("Course / Laboratory", "DevOps & Cloud Automation / Software Engineering Laboratory"),
        ("GitHub Repository", "https://github.com/muawiyah123/Jenkins-Based-Research-Paper-Repository.git"),
        ("Active Branches", "main (Production Baseline) | development (Active Sprint)"),
        ("Test Automation Status", "79/79 Passing Tests (76 Unit/Slice + 3 Selenium E2E Smoke Tests) — 100% Pass")
    ]
    for idx, (label, val) in enumerate(meta_rows):
        r = meta_table.rows[idx]
        c0, c1 = r.cells[0], r.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, HEX_LIGHT_BG if idx < 5 else HEX_SUCCESS_BG)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(label)
        r0.font.name = "Arial"
        r0.font.size = Pt(9)
        r0.font.bold = True
        r0.font.color.rgb = COLOR_NAVY
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.name = "Arial"
        r1.font.size = Pt(9)
        if idx == 5:
            r1.font.bold = True
            r1.font.color.rgb = COLOR_GREEN
        else:
            r1.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Executive Summary
    add_styled_heading(doc, "Executive Summary", 1)
    p_exec = doc.add_paragraph()
    p_exec.paragraph_format.line_spacing = 1.15
    p_exec.paragraph_format.space_after = Pt(6)
    p_exec.add_run(
        "The Jenkins-Based Research Paper Repository is an enterprise-grade academic publishing system engineered "
        "with Spring Boot 3.4.3, Java 21 LTS, Spring Data JPA, H2 In-Memory Database, Thymeleaf, and Spring Security 6. "
        "The platform modernizes academic peer-review workflows by replacing manual, fragmented email submissions with "
        "a role-governed submission and review portal. The software lifecycle is automated through a comprehensive Jenkins CI/CD "
        "pipeline featuring automated checkout, unit and slice testing, artifact packaging, JUnit reporting, containerization "
        "via multi-stage Dockerfile, and deployment orchestration through Ansible."
    )

    # Technology Stack
    add_styled_heading(doc, "1. System Architecture & Technology Stack", 1)
    
    tech_data = [
        ("Component", "Technology", "Version", "Purpose & Architectural Rationale"),
        ("Programming Language", "Java", "21 LTS", "Modern LTS, virtual threads, pattern matching, enhanced security"),
        ("Framework", "Spring Boot", "3.4.3", "Microservice bootstrap, dependency injection, embedded Tomcat"),
        ("Security Layer", "Spring Security", "6.x", "Role-Based Access Control (RBAC), BCrypt hashing, CSRF protection"),
        ("Persistence / ORM", "Spring Data JPA / Hibernate", "6.x", "Declarative repository queries, schema generation, transaction management"),
        ("Database", "H2 In-Memory Database", "2.3.x", "Zero-dependency ephemeral database for deterministic CI test runs"),
        ("Template Engine", "Thymeleaf + Bootstrap 5", "5.3.3", "Server-side rendering, responsive UX, candidate attribution branding"),
        ("CI Automation", "Jenkins CI", "2.580.1", "Declarative multi-stage pipeline, JUnit test archiving, build trend tracking"),
        ("E2E Testing", "Selenium WebDriver", "4.28.1", "Headless Chrome automated browser test suite"),
        ("Containerization", "Docker & Docker Compose", "Multi-stage", "Minimal non-root runtime image, portable container execution"),
        ("Provisioning", "Ansible", "Playbook", "Automated deployment, configuration management, idempotent execution")
    ]
    
    tech_table = doc.add_table(rows=len(tech_data), cols=4)
    tech_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tech_table.autofit = False
    
    col_widths = [Inches(1.5), Inches(1.3), Inches(0.9), Inches(2.8)]
    for r_idx, row in enumerate(tech_data):
        t_row = tech_table.rows[r_idx]
        for c_idx, val in enumerate(row):
            cell = t_row.cells[c_idx]
            cell.width = col_widths[c_idx]
            set_cell_margins(cell, 60, 60, 80, 80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if r_idx == 0:
                set_cell_background(cell, HEX_NAVY)
                r = p.add_run(val)
                r.font.name = "Arial"
                r.font.size = Pt(8.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, HEX_LIGHT_BG if r_idx % 2 == 0 else "FFFFFF")
                r = p.add_run(val)
                r.font.name = "Arial"
                r.font.size = Pt(8)
                r.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 15-Task Implementation Summary
    add_styled_heading(doc, "2. 15-Task Implementation Summary", 1)
    
    tasks_summary = [
        ("Task 01", "Environment & Toolchain Verification", "Java 21, Maven 3.9.9 wrapper, port 8081 allocation", "COMPLETED"),
        ("Task 02", "Spring Boot Scaffolding & Build", "Dependencies, H2 DB configuration, contextLoads test", "COMPLETED"),
        ("Task 03", "Version Control Baseline & Git Init", "Dual-branch model (main/development), conventional commits", "COMPLETED"),
        ("Task 04", "Public Home & Navigation Scaffold", "Bootstrap 5 glassmorphism, candidate badge, quick links", "COMPLETED"),
        ("Task 05", "Domain Model & Persistence", "Paper JPA entity, PaperRepository derived queries, data.sql seed", "COMPLETED"),
        ("Task 06", "Paper Submission Workflow", "Multipart/form-data, Jakarta Bean Validation, flash feedback", "COMPLETED"),
        ("Task 07", "Public Catalog & Keyword Search", "Case-insensitive query search, author email privacy redaction", "COMPLETED"),
        ("Task 08", "Metadata Edit & Resubmission", "Edit workflow restricted to REVISION_REQUESTED lifecycle state", "COMPLETED"),
        ("Task 09", "Spring Security 6 In-Memory RBAC", "RESEARCHER, REVIEWER, ADMIN roles, BCrypt password hashing", "COMPLETED"),
        ("Task 10", "Peer-Review Workflow & State Machine", "5-state finite lifecycle machine, reviewer remarks", "COMPLETED"),
        ("Task 11", "Administrative Metrics Dashboard", "Aggregate status metrics cards, acceptance rate computation", "COMPLETED"),
        ("Task 12", "Input Validation & Error Templates", "Custom 403, 404, 500 HTML error pages, 76 test assertions", "COMPLETED"),
        ("Task 13", "Jenkins CI Declarative Pipeline", "6-stage automated pipeline, JUnit XML report publisher", "COMPLETED"),
        ("Task 14", "Automated Selenium E2E Smoke Tests", "Headless Chrome browser tests, DOM validation, CI ready", "COMPLETED"),
        ("Task 15", "Docker Packaging & Ansible Automation", "Multi-stage Dockerfile, docker-compose.yml, ansible playbook", "COMPLETED")
    ]
    
    sum_table = doc.add_table(rows=len(tasks_summary) + 1, cols=4)
    sum_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sum_table.autofit = False
    
    sum_widths = [Inches(1.0), Inches(2.3), Inches(2.4), Inches(0.8)]
    headers = ["Task ID", "Task Title", "Technical Scope & Deliverables", "Status"]
    
    for c_idx, h in enumerate(headers):
        cell = sum_table.rows[0].cells[c_idx]
        cell.width = sum_widths[c_idx]
        set_cell_margins(cell, 60, 60, 80, 80)
        set_cell_background(cell, HEX_NAVY)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = "Arial"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    for r_idx, row in enumerate(tasks_summary):
        t_row = sum_table.rows[r_idx + 1]
        for c_idx, val in enumerate(row):
            cell = t_row.cells[c_idx]
            cell.width = sum_widths[c_idx]
            set_cell_margins(cell, 50, 50, 80, 80)
            set_cell_background(cell, HEX_LIGHT_BG if r_idx % 2 == 0 else "FFFFFF")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(val)
            r.font.name = "Arial"
            r.font.size = Pt(8)
            if c_idx == 3:
                r.font.bold = True
                r.font.color.rgb = COLOR_GREEN
            else:
                r.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Visual Proof of Work Figures
    add_styled_heading(doc, "3. Visual Proof of Work & Empirical Evidence Dossier", 1)
    
    p_ev = doc.add_paragraph()
    p_ev.paragraph_format.line_spacing = 1.15
    p_ev.add_run(
        "The following 15 figures represent empirical proof of work across all terminal commands, "
        "Jenkins CI execution stages, and Spring Boot web application user journeys. All terminal captures "
        "reflect native Windows 11 Windows Terminal and PowerShell execution."
    )

    # Figures List
    figures = [
        ("cmd-01-mvnw-test.png", "1", "PowerShell Terminal Execution of Standard Maven Test Suite (76/76 Tests Passed, BUILD SUCCESS)"),
        ("cmd-02-selenium-e2e.png", "2", "PowerShell Terminal Execution of Headless Chrome Selenium E2E Test Suite (3/3 Tests Passed)"),
        ("cmd-03-git-log.png", "3", "PowerShell Terminal Execution of Git Branching & Conventional Commit Log"),
        ("docker-06-cli-build-verification.png", "4", "Terminal Execution of Docker Container Packaging Verification"),
        ("jenkins-01-dashboard-welcome.png", "5", "Jenkins Automation Controller Running Locally on Port 8080"),
        ("jenkins-02-pipeline-configuration.png", "6", "Jenkins Pipeline Configuration Targeting Git SCM and Jenkinsfile"),
        ("jenkins-03-branch-specifier.png", "7", "Jenkins SCM Branch Specifier Configuration (*/main)"),
        ("jenkins-04-build1-success-details.png", "8", "Jenkins Build #1 Detailed Execution View — 100% Successful Build & Packaged JAR Artifact"),
        ("jenkins-05-build-trend-test-results.png", "9", "Jenkins Test Result Trend Graph Across Consecutive Automated Pipeline Runs"),
        ("ui-01-homepage.png", "10", "Public Home Portal with Responsive Navigation and Candidate Attribution Banner"),
        ("ui-02-catalog.png", "11", "Public Research Paper Catalog with Real-Time Search and Author Email Privacy Redaction"),
        ("ui-03-login.png", "12", "Spring Security 6 Custom Login Screen with Academic Role Credentials"),
        ("ui-04-submission-form.png", "13", "Researcher Paper Submission Form with Jakarta Bean Validation Rules"),
        ("ui-05-review-queue.png", "14", "Reviewer Peer-Review Evaluation Portal and 5-State Finite Lifecycle Machine"),
        ("ui-06-admin-dashboard.png", "15", "Administrative Metrics Dashboard with Live Status Counts and Acceptance Rate Calculation")
    ]

    for img_name, fig_num, caption in figures:
        add_figure(doc, img_name, fig_num, caption)

    # Quality Assurance & Test Pyramid Matrix
    add_styled_heading(doc, "4. Quality Assurance & Test Automation Matrix", 1)
    
    qa_data = [
        ("Testing Slice", "Class / Specification", "Test Count", "Execution Scope & Verification"),
        ("Context / Scaffolding", "RepositoryApplicationTests", "1", "Spring Boot context bootstrapping and H2 connection verification"),
        ("JPA Data Access Layer", "PaperRepositoryTest", "5", "CRUD operations, derived keyword search, and status count queries"),
        ("Service Layer", "PaperServiceTest", "25", "5-state review lifecycle state machine transitions and business validation"),
        ("Web MVC / Home", "HomeControllerTest", "3", "Public homepage routing, model attributes, and candidate badge render"),
        ("Web MVC / Submissions", "PaperControllerTest", "15", "Researcher submission validation, CSRF handling, and edit workflows"),
        ("Web MVC / Review", "ReviewControllerTest", "10", "Reviewer queue authorization, status transitions, and review feedback"),
        ("Web MVC / Dashboard", "DashboardControllerTest", "5", "Admin dashboard aggregate calculations and view model population"),
        ("Spring Security RBAC", "SecurityConfigTest", "12", "URL authorization rules, unauthenticated redirects, role enforcement"),
        ("End-to-End (E2E) Smoke", "PaperSubmissionUiTest", "3", "Headless Chrome browser tests, search queries, email redaction"),
        ("Total Quality Assurance", "All Architectural Slices", "79", "100% Pass Rate across Unit, Slice, Integration, and E2E Tests")
    ]
    
    qa_table = doc.add_table(rows=len(qa_data), cols=4)
    qa_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    qa_table.autofit = False
    
    qa_widths = [Inches(1.5), Inches(1.8), Inches(0.8), Inches(2.4)]
    for r_idx, row in enumerate(qa_data):
        t_row = qa_table.rows[r_idx]
        for c_idx, val in enumerate(row):
            cell = t_row.cells[c_idx]
            cell.width = qa_widths[c_idx]
            set_cell_margins(cell, 60, 60, 80, 80)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            if r_idx == 0:
                set_cell_background(cell, HEX_NAVY)
                r = p.add_run(val)
                r.font.name = "Arial"
                r.font.size = Pt(8.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
            elif r_idx == len(qa_data) - 1:
                set_cell_background(cell, HEX_SUCCESS_BG)
                r = p.add_run(val)
                r.font.name = "Arial"
                r.font.size = Pt(8.5)
                r.font.bold = True
                r.font.color.rgb = COLOR_GREEN
            else:
                set_cell_background(cell, HEX_LIGHT_BG if r_idx % 2 == 0 else "FFFFFF")
                r = p.add_run(val)
                r.font.name = "Arial"
                r.font.size = Pt(8)
                r.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Conclusion & Sign-off
    add_styled_heading(doc, "5. Conclusion & Academic Certification", 1)
    p_conc = doc.add_paragraph()
    p_conc.paragraph_format.line_spacing = 1.15
    p_conc.add_run(
        "The Jenkins-Based Research Paper Repository successfully fulfills all functional, architectural, "
        "and continuous integration requirements established for the academic semester. The enterprise Spring Boot application "
        "demonstrates robust role-based security, rigorous validation, and complete test automation. The automated Jenkins CI "
        "pipeline, multi-stage Docker containerization, and Ansible deployment orchestration validate the engineering rigor "
        "of the software lifecycle from source control to deployment."
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Sign-off Box
    sign_table = doc.add_table(rows=1, cols=1)
    sign_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    sign_table.autofit = False
    sc = sign_table.rows[0].cells[0]
    sc.width = Inches(6.5)
    set_cell_margins(sc, 120, 120, 160, 160)
    set_cell_background(sc, "F8FAFC")
    
    sp = sc.paragraphs[0]
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp.paragraph_format.space_after = Pt(0)
    s1 = sp.add_run("Project Implementation & Documentation Certified by:\n")
    s1.font.name = "Arial"
    s1.font.size = Pt(9.5)
    s1.font.italic = True
    s1.font.color.rgb = COLOR_GRAY
    
    s2 = sp.add_run("MUAWIYAH SAHIL NARVEL | ROLL NO: 23108A0020\n")
    s2.font.name = "Arial"
    s2.font.size = Pt(11)
    s2.font.bold = True
    s2.font.color.rgb = COLOR_NAVY
    
    s3 = sp.add_run("Department of Computer Engineering / Information Technology\nDate of Verification: October 15, 2026")
    s3.font.name = "Arial"
    s3.font.size = Pt(9)
    s3.font.color.rgb = COLOR_GRAY

    doc.save(out_docx)
    print(f"Generated Master Word Report: {out_docx}")

if __name__ == "__main__":
    main()
